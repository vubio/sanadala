"""Build separate grammar lesson datasets from seven HSK workbooks.

Usage:
    python scripts/convert_grammar.py /path/to/folder-containing-Grammar_HSK_*.xlsx

Publishing includes learner-facing text from Lessons, Lesson_Items, Teaching_Items,
Examples and Grammar_Nodes. Editorial logs, source lines and original XLSX files
are not published.
"""

from collections import defaultdict
import json
from pathlib import Path
import sys

from openpyxl import load_workbook


LEVELS = [('1', '1'), ('2', '2'), ('3', '3'), ('4', '4'),
          ('5', '5'), ('6', '6'), ('79', '7_9')]
TOPIC_TRANSLATIONS = {
    '词类': 'Từ loại', '句子的类型': 'Kiểu câu', '短语': 'Cụm từ',
    '强调的方法': 'Cách nhấn mạnh', '口语格式': 'Cấu trúc khẩu ngữ',
    '句子成分': 'Thành phần câu', '固定格式': 'Mẫu cố định',
    '复句': 'Câu phức', '介词': 'Giới từ', '提问的方法': 'Cách hỏi',
    '副词': 'Phó từ', '补语': 'Bổ ngữ', '特殊表达法': 'Cách diễn đạt đặc biệt',
    '特殊句型': 'Mẫu câu đặc biệt', '句群': 'Đoạn văn',
    '动作的态': 'Thể của hành động', '动词': 'Động từ',
    '固定短语': 'Cụm từ cố định', '量词': 'Lượng từ',
    '代词': 'Đại từ', '结构类型': 'Loại cấu trúc',
    '语素': 'Hình vị', '前缀': 'Tiền tố', '后缀': 'Hậu tố',
    '连词': 'Liên từ', '拟声词': 'Từ tượng thanh',
    '主语': 'Chủ ngữ', '宾语': 'Tân ngữ', '定语': 'Định ngữ',
    '句型': 'Mẫu câu',
}
OUTPUT_DIR = Path(__file__).resolve().parents[1] / 'data'


def records(workbook, name):
    iterator = workbook[name].iter_rows(values_only=True)
    headers = next(iterator)
    return [dict(zip(headers, row)) for row in iterator if any(x is not None for x in row)]


def enabled(row):
    return row.get('Enabled') not in (False, 0, 'FALSE')


def text(value):
    return str(value).strip() if value is not None else ''


def get_topic(point_id, nodes):
    node = nodes.get(point_id)
    ancestors = []
    seen = set()
    while node and node.get('Parent_ID') and node['Node_ID'] not in seen:
        seen.add(node['Node_ID'])
        node = nodes.get(node['Parent_ID'])
        if node and 'Category' in text(node.get('Node_Type')):
            title = text(node.get('Title_Chinese'))
            # Some HSK 3 headings are root categories (e.g., 词类) rather than
            # a generic level root; retain those as real study topics.
            is_level_root = not node.get('Parent_ID') and '级语法点' in title
            if not is_level_root:
                ancestors.insert(0, node)
    if not ancestors:
        return 'Cấu trúc và cách dùng'
    first = ancestors[0]
    return (text(first.get('Title_Vietnamese')) or
            TOPIC_TRANSLATIONS.get(text(first.get('Title_Chinese')),
                                   text(first.get('Title_Chinese'))))


def convert_level(source, key):
    workbook = load_workbook(source, read_only=True, data_only=True)
    lessons = [r for r in records(workbook, 'Lessons') if enabled(r)]
    teaching = {r['Item_ID']: r for r in records(workbook, 'Teaching_Items') if enabled(r)}
    lesson_links = defaultdict(list)
    for r in records(workbook, 'Lesson_Items'):
        lesson_links[r['Lesson_ID']].append(r)
    examples = defaultdict(list)
    for r in records(workbook, 'Examples'):
        if enabled(r):
            examples[r['Item_ID']].append(r)
    nodes = {r['Node_ID']: r for r in records(workbook, 'Grammar_Nodes')}

    output = []
    used_items = set()
    used_examples = set()
    for lesson in sorted(lessons, key=lambda r: int(r['Lesson_Order'])):
        lesson_id = lesson['Lesson_ID']
        parts = []
        for link in sorted(lesson_links[lesson_id], key=lambda r: int(r['Position'])):
            source_item = teaching.get(link['Item_ID'])
            if not source_item:
                raise ValueError(f'{source}: missing teaching item {link["Item_ID"]}')
            used_items.add(source_item['Item_ID'])
            sentences = []
            for ex in sorted(examples[source_item['Item_ID']],
                             key=lambda r: (int(r['Example_Order'] or 0),
                                            int(r['Group_Order'] or 0))):
                used_examples.add(ex['Example_ID'])
                if not all(text(ex.get(f)) for f in
                           ('Example_Chinese', 'Example_Pinyin', 'Example_Vietnamese')):
                    raise ValueError(f'{source}: missing example content {ex["Example_ID"]}')
                sentences.append({
                    'id': ex['Example_ID'],
                    'cn': text(ex['Example_Chinese']),
                    'pinyin': text(ex['Example_Pinyin']),
                    'sino': text(ex.get('Example_Sinovietnamese')),
                    'vi': text(ex['Example_Vietnamese']),
                    'type': text(ex.get('Example_Type')).lower(),
                    'group': text(ex.get('Group_ID')),
                    'role': text(ex.get('Example_Role')),
                    'speaker': text(ex.get('Speaker')),
                })
            if not sentences:
                raise ValueError(f'{source}: no examples for {source_item["Item_ID"]}')
            parts.append({
                'id': source_item['Item_ID'],
                'focus': text(source_item.get('Focus_Chinese')),
                'pattern': text(source_item.get('Pattern_Chinese')),
                'explanation': text(source_item.get('Explanation_Vietnamese')),
                'spoken': text(source_item.get('Focus_Spoken_Chinese')),
                'examples': sentences,
            })
        if not parts:
            raise ValueError(f'{source}: no teaching items for {lesson_id}')
        output.append({
            'id': lesson_id,
            'number': int(lesson['Lesson_Order']),
            'title': text(lesson['Title_Vietnamese']),
            'intro': text(lesson.get('Intro_Vietnamese')),
            'topic': get_topic(lesson['Point_ID'], nodes),
            'items': parts,
        })

    leftover_items = set(teaching) - used_items
    leftover_examples = {ex['Example_ID'] for rows in examples.values() for ex in rows} - used_examples
    if leftover_items or leftover_examples:
        raise ValueError(f'{source}: unlinked items {leftover_items} / examples {leftover_examples}')
    output_file = f'grammar-{key}.json'
    (OUTPUT_DIR / output_file).write_text(
        json.dumps(output, ensure_ascii=False, separators=(',', ':')) + '\n',
        encoding='utf-8')
    return {'key': key, 'label': 'HSK 7–9' if key == '79' else f'HSK {key}',
            'file': output_file, 'lessons': len(output),
            'items': len(used_items), 'examples': len(used_examples)}


def main(folder):
    catalog = {'version': 1, 'levels': []}
    for key, suffix in LEVELS:
        filename = folder / f'Grammar_HSK_{suffix}.xlsx'
        if not filename.is_file():
            raise FileNotFoundError(filename)
        stats = convert_level(filename, key)
        catalog['levels'].append(stats)
        print(f'{stats["label"]}: {stats["lessons"]} lessons, '
              f'{stats["items"]} items, {stats["examples"]} examples')
    for prop in ('lessons', 'items', 'examples'):
        catalog['total' + prop.title()] = sum(x[prop] for x in catalog['levels'])
    (OUTPUT_DIR / 'grammar-catalog.json').write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        raise SystemExit('Usage: python scripts/convert_grammar.py <folder-with-workbooks>')
    main(Path(sys.argv[1]))
