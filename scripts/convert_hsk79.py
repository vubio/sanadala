"""Convert the separately maintained combined HSK 7–9 workbook to site JSON.

Usage: python scripts/convert_hsk79.py /path/to/HSK7_Examples_1-5636_Complete_with_Example_Readings.xlsx

Only the vocabulary sheet is published. The source workbook and its editorial
Gloss_Changes tab remain private; the public JSON contains learning content.
"""

import json
from pathlib import Path
import sys

from openpyxl import load_workbook


EXPECTED_HEADERS = (
    'Level', 'Number', 'Vocabulary', 'Sinovietnamese', 'Pinyin',
    'Vietnamese', 'Example_Chinese', 'Example_Sinovietnamese',
    'Example_Pinyin', 'Example_Vietnamese',
)
CHUNK_SIZE = 500
OUTPUT_DIR = Path(__file__).resolve().parents[1] / 'data'


def main(source_path):
    book = load_workbook(source_path, read_only=True, data_only=True)
    sheet = book['HSK_1-5636']
    rows = sheet.iter_rows(values_only=True)
    headers = tuple(next(rows))
    if headers != EXPECTED_HEADERS:
        raise ValueError(f'Unexpected headers: {headers!r}')

    converted = []
    for source in rows:
        if all(value is None for value in source):
            continue
        level, number, word, sino, pinyin, vietnamese, example, example_sino, example_pinyin, example_vi = source
        if level != 'HSK 7-9':
            raise ValueError(f'Unexpected combined-level label: {level!r} at {number}')
        if number != len(converted) + 1:
            raise ValueError(f'Out-of-order number: {number} after {len(converted)} rows')
        if any(value is None or str(value).strip() == '' for value in source):
            raise ValueError(f'Missing field in row {number}')
        # HSK 3.0 site's 11-field row layout; the source has no Traditional or Note.
        converted.append([number, word, sino, pinyin, vietnamese,
                          example, '', example_pinyin, example_sino, example_vi, ''])

    if len(converted) != 5636:
        raise ValueError(f'Expected 5,636 entries; found {len(converted)}')
    chunks = []
    for start in range(0, len(converted), CHUNK_SIZE):
        name = f'hsk79-{start // CHUNK_SIZE + 1}.json'
        (OUTPUT_DIR / name).write_text(
            json.dumps(converted[start:start + CHUNK_SIZE], ensure_ascii=False,
                       separators=(',', ':')) + '\n', encoding='utf-8')
        chunks.append(name)

    catalog = {
        'version': 1,
        'edition': 'HSK 3.0 · Cấp 7–9',
        'total': len(converted),
        'levels': [{'level': 7, 'label': 'HSK 7–9',
                    'files': chunks, 'count': len(converted)}],
    }
    (OUTPUT_DIR / 'catalog79.json').write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Converted {len(converted)} entries into {len(chunks)} JSON chunks')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        raise SystemExit('Usage: python scripts/convert_hsk79.py <HSK7-9.xlsx>')
    main(Path(sys.argv[1]))
